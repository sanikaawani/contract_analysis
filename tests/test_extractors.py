"""
Unit tests for text extractors and document parsers.
"""

import os
import io
import pytest
from docx import Document
from src.extractors import extract_text_from_file, clean_contract_text


def test_clean_contract_text():
    raw = "Section 1.1\u201cServices\u201d\n\n\n\nPage 1 of 5\n\nVendor shall deliver."
    cleaned = clean_contract_text(raw)
    assert 'Page 1 of 5' not in cleaned
    assert '"Services"' in cleaned
    assert 'Section 1.1' in cleaned


def test_extract_from_txt_bytes():
    text_content = "MASTER SERVICE AGREEMENT\n1. Payment Terms\nClient shall pay in 30 days."
    res = extract_text_from_file(text_content.encode('utf-8'), filename="agreement.txt")
    assert res['word_count'] > 5
    assert res['detected_title'] == "MASTER SERVICE AGREEMENT"


def test_extract_from_docx_buffer():
    doc = Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=1)
    doc.add_paragraph("This NDA binds both parties for 2 years.")
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = extract_text_from_file(docx_bytes, filename="nda.docx")
    assert "NON-DISCLOSURE AGREEMENT" in res['cleaned_text']
    assert "2 years" in res['cleaned_text']
